"""core.py - keyword ATS score, targeted line suggestions, and in-place edits.

PDF  -> edited directly (no conversion, layout stays identical)
DOCX -> edited directly with python-docx (LibreOffice only used to also export a PDF)
"""
import os, io, re, json, unicodedata, hashlib, shutil, subprocess, tempfile, base64
from difflib import SequenceMatcher
import pymupdf as fitz
from groq import Groq
from docx import Document

MODEL = os.getenv("GROQ_MODEL", "openai/gpt-oss-120b")
MAX_BYTES = 5 * 1024 * 1024
_kw_cache: dict[str, list[str]] = {}
_norm = lambda s: re.sub(r"[\W_]+", " ", s.lower()).strip()
BULLETS = set("•-–—▪●◦*·■")
GLYPH = re.compile(r"^[^\w(]+")     # leading bullet symbols


# ── Reading resumes ─────────────────────────────────────────────
def _kind(data: bytes):
    return "pdf" if data[:4] == b"%PDF" else "docx" if data[:2] == b"PK" else None


def pdf_text(data: bytes) -> str:
    with fitz.open(stream=data, filetype="pdf") as doc:
        return unicodedata.normalize("NFKC", "\n".join(page.get_text() for page in doc)).strip()


def _iter_paras(container, seen):
    for p in container.paragraphs:
        if id(p._p) not in seen:
            seen.add(id(p._p)); yield p
    for t in getattr(container, "tables", []):
        for row in t.rows:
            for cell in row.cells:
                yield from _iter_paras(cell, seen)


def resume_text(data: bytes) -> str:
    k = _kind(data)
    if k == "pdf":
        return pdf_text(data)
    if k == "docx":
        return "\n".join(p.text for p in _iter_paras(Document(io.BytesIO(data)), set())).strip()
    raise ValueError("Upload a PDF or Word (.docx) resume.")


# ── LLM helpers ─────────────────────────────────────────────────
def _client(key: str | None) -> Groq:
    key = key or os.getenv("GROQ_API_KEY")
    if not key:
        raise ValueError("No Groq API key. Add one in the app or set GROQ_API_KEY on the server.")
    return Groq(api_key=key)


def _json_call(client, system: str, user: str, max_tokens: int) -> dict:
    extra = {}
    if "gpt-oss" in MODEL:
        extra["reasoning_effort"] = "low"          # reasoning models spend tokens thinking
        max_tokens = max_tokens * 2 + 2000
    r = client.chat.completions.create(
        model=MODEL,
        messages=[{"role": "system", "content": system}, {"role": "user", "content": user}],
        response_format={"type": "json_object"},
        temperature=0.2,
        max_tokens=max_tokens,
        **extra,
    )
    return json.loads(r.choices[0].message.content)


# ── ATS score: deterministic keyword match ──────────────────────
def get_keywords(jd: str, key: str | None) -> list[str]:
    """Hard skills/tools from the JD, cached per JD so before/after scores use the same list."""
    h = hashlib.sha1(jd.strip().encode()).hexdigest()
    if h not in _kw_cache:
        data = _json_call(
            _client(key),
            "You extract ATS keywords from job descriptions. Reply with JSON only.",
            'List the concrete hard skills, tools, technologies, methods and certifications in this job '
            'description, worded as they would appear on a resume (1-3 words each, max 25, no soft skills, '
            'no duplicates). Format: {"keywords": ["Python", "SQL"]}\n\nJOB DESCRIPTION:\n' + jd[:4000],
            600,
        )
        seen, out = set(), []
        for k in data.get("keywords", []):
            if isinstance(k, str) and k.strip() and k.lower() not in seen:
                seen.add(k.lower()); out.append(k.strip())
        _kw_cache[h] = out[:25]
        if len(_kw_cache) > 200:
            _kw_cache.pop(next(iter(_kw_cache)))
    return _kw_cache[h]


def _has(text_l: str, kw: str) -> bool:
    return re.search(r"(?<![a-z0-9])" + re.escape(kw.lower()) + r"(?![a-z0-9])", text_l) is not None


def score(text: str, kws: list[str]):
    t = text.lower()
    matched = [k for k in kws if _has(t, k)]
    missing = [k for k in kws if k not in matched]
    pct = round(100 * len(matched) / len(kws)) if kws else 0
    return pct, matched, missing


# ── Analyze: score + line-level suggestions + questions ─────────
def _on_resume(original: str, text_norm: str) -> bool:
    """True if `original` really appears in the resume (guards against invented lines)."""
    o = _norm(original)
    if len(o) < 20:
        return False
    m = SequenceMatcher(None, o, text_norm, autojunk=False).find_longest_match(0, len(o), 0, len(text_norm))
    return m.size / len(o) >= 0.6


def analyze(data: bytes, jd: str, key: str | None) -> dict:
    text = resume_text(data)
    if len(text) < 50:
        raise ValueError("No readable text in this file. Scanned or image-only resumes are not supported.")
    kws = get_keywords(jd, key)
    pct, matched, missing = score(text, kws)
    res = _json_call(
        _client(key),
        "You are a senior technical recruiter. Reply with JSON only.",
        f"Keywords missing from the resume: {missing}\n\n"
        'Return {"improvements": [{"original","suggested","reason"}], "questions": [...]}.\n'
        "- improvements: identify ALL meaningful single-line replacements needed in this resume for this job. "
        "Do not stop at 6. `original` must be ONE existing bullet or line copied exactly from the resume. "
        "`suggested` may reword only that line to naturally use a missing keyword or the job's wording, and only "
        "if the existing line already demonstrates that skill. Never invent tools, experience, responsibilities, "
        "certifications, metrics or numbers. Do not make cosmetic rewrites. Do not suggest a change merely because "
        "different wording sounds better. Keep each suggested line close to the original length and formatting. "
        "Every missing keyword that can be honestly supported by an existing resume line should be addressed in this "
        "single response. If a missing keyword cannot be supported by the resume, do not force it into a line. "
        "If no honest improvement is possible, return an empty list.\n"
        "- Before returning improvements, check the entire resume against ALL missing keywords and do not leave an "
        "obvious supported keyword for a later analysis.\n"
        "- questions: exactly 5 interview questions tailored to this job description and this candidate's projects.\n\n"
        f"RESUME:\n{text[:6000]}\n\nJOB DESCRIPTION:\n{jd[:3000]}",
        4000,
    )
    tn = _norm(text)
    edits = []
    for e in res.get("improvements", []):
        if isinstance(e, dict) and all(isinstance(e.get(k), str) for k in ("original", "suggested")):
            if e["suggested"].strip() and _on_resume(e["original"], tn):
                edits.append({"original": e["original"].strip(), "suggested": e["suggested"].strip(),
                              "reason": str(e.get("reason", ""))[:120]})
    qs = [s.strip() for s in res.get("questions", []) if isinstance(s, str) and s.strip()]
    return {"score": pct, "matched": matched, "missing": missing, "improvements": edits, "questions": qs}


# ── Shared: decide which paragraph/block each edit targets ──────
def _match_edits(edits, texts):
    used, plan, skipped = set(), {}, []
    for e in edits:
        o = _norm(e["original"]); best, top = None, 0.0
        for i, t in enumerate(texts):
            if i in used:
                continue
            r = SequenceMatcher(None, o, _norm(t)).ratio()
            if r > top:
                best, top = i, r
        if best is not None and top >= 0.6:
            used.add(best); plan[best] = e["suggested"]
        else:
            skipped.append(e["original"])
    return plan, skipped


def _skills_line(texts, matched, taken):
    cands = [i for i, t in enumerate(texts) if t.count(",") >= 2 and i not in taken]
    hits = lambda i: sum(_has(texts[i].lower(), k) for k in matched)
    best = max(cands, key=hits, default=None)
    return best if best is not None and hits(best) > 0 else None


# ── PDF: edit in place (cover the old line, write the new text at the same baseline) ──
_LATIN = str.maketrans({"\u2019": "'", "\u2018": "'", "\u201c": '"', "\u201d": '"', "\u2013": "-", "\u2014": "-",
                        "\u2022": "-", "\u00a0": " ", "\u2192": "->"})
_latin = lambda t: t.translate(_LATIN).encode("latin-1", "ignore").decode("latin-1")
_rgb = lambda c: ((c >> 16 & 255) / 255, (c >> 8 & 255) / 255, (c & 255) / 255)


def _fonts(flags):
    fam = "cour" if flags & 8 else "tiro" if flags & 4 else "helv"
    return {"helv": ("helv", "hebo"), "tiro": ("tiro", "tibo"), "cour": ("cour", "cobo")}[fam]


def _geom(b):
    lines = [[s for s in l["spans"] if s["text"].strip()] for l in b["lines"]]
    lines = [l for l in lines if l]
    if lines and len(lines[0]) and lines[0][0]["text"].strip() in BULLETS:
        lines[0] = lines[0][1:]                        # leave a separate bullet glyph untouched
        lines = [l for l in lines if l]
    spans = [s for l in lines for s in l]
    if not spans:
        return None
    first, body = spans[0], spans[-1]
    lead = GLYPH.match(first["text"])
    label = first["text"].strip() if (first["flags"] & 16 and first["text"].strip().endswith(":") and len(spans) > 1) else None
    pitch = (lines[1][0]["origin"][1] - lines[0][0]["origin"][1]) if len(lines) > 1 else None
    return dict(
        rect=fitz.Rect(min(s["bbox"][0] for s in spans), min(s["bbox"][1] for s in spans),
                       max(s["bbox"][2] for s in spans), max(s["bbox"][3] for s in spans)),
        x0=min(s["origin"][0] for s in spans), y0=lines[0][0]["origin"][1], n=len(lines), pitch=pitch,
        size=body["size"], color=_rgb(body["color"]), fonts=_fonts(body["flags"]),
        lead=lead.group().strip() if lead else "", label=label)


def _write(page, g, new):
    reg, bold = g["fonts"]
    new = _latin(GLYPH.sub("", new.strip()))
    if g["lead"]:
        new = g["lead"] + " " + new
    label = _latin(g["label"]) if g["label"] else None
    toks = [(w, bool(label) and i < len(label.split())) for i, w in enumerate(new.split())]
    if not label or not new.startswith(label):
        toks = [(w, False) for w, _ in toks]
    limit = min(g["rect"].x1 - g["x0"] + 12, page.rect.width - 36 - g["x0"])
    for scale in (1.0, 0.95, 0.9, 0.85):               # shrink slightly before adding a line
        size = g["size"] * scale
        L = lambda t, b=False: fitz.get_text_length(t, fontname=bold if b else reg, fontsize=size)
        lines, cur = [[]], 0.0
        for w, b in toks:
            sp = L(" ") if lines[-1] else 0
            if lines[-1] and cur + sp + L(w, b) > limit:
                lines.append([]); cur, sp = 0.0, 0
            lines[-1].append((w, b)); cur += sp + L(w, b)
        if len(lines) <= g["n"]:
            break
    pitch = g["pitch"] or size * 1.2
    for k, line in enumerate(lines):
        segs = []
        for w, b in line:
            if segs and segs[-1][1] == b:
                segs[-1][0] += " " + w
            else:
                segs.append([w, b])
        x, y = g["x0"], g["y0"] + k * pitch
        for text, b in segs:
            page.insert_text((x, y), text, fontname=bold if b else reg, fontsize=size, color=g["color"])
            x += L(text, b) + L(" ")


def _append_to_line(page, blocks_on_page, skip_ids, matched, words):
    """Add skills at the end of the best skills line (the line with most matched keywords)."""
    best, top = None, 0
    for bi, b in blocks_on_page:
        if bi in skip_ids:
            continue
        for l in b["lines"]:
            t = "".join(s["text"] for s in l["spans"])
            h = sum(_has(t.lower(), k) for k in matched)
            if t.count(",") >= 2 and h > top:
                best, top = l, h
    if best is None:
        return False
    sp = [s for s in best["spans"] if s["text"].strip()][-1]
    reg, _ = _fonts(sp["flags"])
    txt = _latin(", " + ", ".join(words))
    end = sp["bbox"][2]
    if end + fitz.get_text_length(txt, fontname=reg, fontsize=sp["size"]) > page.rect.width - 30:
        return False                                   # would run off the page
    page.insert_text((end, sp["origin"][1]), txt, fontname=reg, fontsize=sp["size"], color=_rgb(sp["color"]))
    return True


def _edit_pdf(data, edits, new_skills, matched):
    doc = fitz.open(stream=data, filetype="pdf")
    blocks, texts = [], []
    for pno, page in enumerate(doc):
        for b in page.get_text("dict")["blocks"]:
            if b["type"] != 0:
                continue
            t = " ".join("".join(s["text"] for s in l["spans"]) for l in b["lines"]).strip()
            if len(t) >= 10 and _geom(b):
                blocks.append((pno, b)); texts.append(t)

    plan, skipped = _match_edits(edits, texts)
    by_page = {}
    for i, new in plan.items():
        by_page.setdefault(blocks[i][0], []).append((_geom(blocks[i][1]), new))
    for pno, items in by_page.items():
        page = doc[pno]
        for g, _ in items:
            page.add_redact_annot(g["rect"] + (-1, 0, 1, 0), fill=(1, 1, 1))
        page.apply_redactions(images=0)                # never touch images
        for g, new in items:
            _write(page, g, new)

    skills_ok = None
    if new_skills:
        skills_ok = False
        for pno, page in enumerate(doc):
            on_page = [(i, blocks[i][1]) for i in range(len(blocks)) if blocks[i][0] == pno]
            if _append_to_line(page, on_page, set(plan), matched, new_skills):
                skills_ok = True
                break
    out = doc.tobytes(garbage=3, deflate=True)
    doc.close()
    return out, len(plan), skipped, skills_ok


# ── DOCX: edit paragraphs directly ──────────────────────────────
def _set_text(p, new: str):
    """Replace paragraph text but keep the first run's formatting (and a bold 'Label:' run)."""
    old, runs = p.text, p.runs
    new = GLYPH.sub("", new.strip())
    if not runs or not new:
        return
    k = next((i for i, r in enumerate(runs) if r.text.strip().endswith(":")), None)
    if k is not None and k < len(runs) - 1 and ":" in new:
        j = new.index(":")
        lead = GLYPH.match(runs[k].text)
        runs[k].text = (lead.group() if lead else "") + new[: j + 1] + " "
        runs[k + 1].text = new[j + 1:].strip()
        for r in runs[k + 2:]:
            r.text = ""
    else:
        m = GLYPH.match(old)
        runs[0].text = (m.group() if m else "") + new
        for r in runs[1:]:
            r.text = ""


def _edit_docx(data, edits, new_skills, matched):
    doc = Document(io.BytesIO(data))
    paras = [p for p in _iter_paras(doc, set()) if len(p.text.strip()) >= 10]
    texts = [p.text for p in paras]
    plan, skipped = _match_edits(edits, texts)
    for i, new in plan.items():
        _set_text(paras[i], new)
    skills_ok = None
    if new_skills:
        j = _skills_line(texts, matched, set(plan))
        skills_ok = j is not None and bool(paras[j].runs)
        if skills_ok:
            last = paras[j].runs[-1]
            last.text = last.text.rstrip().rstrip(".") + ", " + ", ".join(new_skills)
    buf = io.BytesIO(); doc.save(buf)
    return buf.getvalue(), len(plan), skipped, skills_ok


def _soffice():
    for c in (os.getenv("SOFFICE_PATH"), shutil.which("soffice"), shutil.which("libreoffice"),
              r"C:\Program Files\LibreOffice\program\soffice.exe",
              r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
              "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if c and os.path.exists(c):
            return c
    return None


def _docx_to_pdf(docx_bytes: bytes):
    """Optional PDF export for .docx uploads. Returns None if LibreOffice is unavailable."""
    exe = _soffice()
    if not exe:
        return None
    try:
        with tempfile.TemporaryDirectory() as d:
            p = os.path.join(d, "r.docx"); open(p, "wb").write(docx_bytes)
            subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", d, p],
                           check=True, timeout=120, capture_output=True)
            return open(os.path.join(d, "r.pdf"), "rb").read()
    except Exception:
        return None


# ── Improve ─────────────────────────────────────────────────────
def improve(data: bytes, jd: str, edits: list[dict], extra_skills: list[str], key: str | None) -> dict:
    kind = _kind(data)
    before_text = resume_text(data)                    # also validates the file type
    kws = get_keywords(jd, key)
    before, matched, _ = score(before_text, kws)
    new_skills = [s for s in extra_skills if not _has(before_text.lower(), s)]

    if kind == "pdf":
        out_pdf, applied, skipped, skills_ok = _edit_pdf(data, edits, new_skills, matched)
        out_docx = None
    else:
        out_docx, applied, skipped, skills_ok = _edit_docx(data, edits, new_skills, matched)
        out_pdf = _docx_to_pdf(out_docx)

    after, _, still_missing = score(resume_text(out_pdf or out_docx), kws)
    b64 = lambda b: base64.b64encode(b).decode() if b else None
    return {
        "score_before": before, "score_after": after, "missing_after": still_missing,
        "applied": applied, "skipped": skipped, "skills_added": skills_ok,
        "pdf": b64(out_pdf), "docx": b64(out_docx),
    }